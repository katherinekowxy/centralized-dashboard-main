"""Electrical Excel ingest: strict checks, preview then commit, the batch record, and idempotent re-import."""

import hashlib
import io
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from demo import ingest_excel, site_db, site_seed

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "examples" / "site"
SAMPLE = SAMPLES / "electrical-sample.xlsx"
BAD_SAMPLE = SAMPLES / "electrical-bad-sample.xlsx"
AS_OF = datetime(2026, 10, 12, 6, 0, tzinfo=timezone.utc)  # the seed's 7-day history covers the samples' 09 Oct window
NOW = AS_OF + timedelta(minutes=7)                           # 14:07 SGT, so batch IDs carry 20261012
HEADER = ("meter_ref", "timestamp", "point", "value", "unit")
GOOD = ("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", 230.4, "V")
OLD = "2026-09-01T08:00:00+08:00"                            # outside the seed's history: a brand-new reading


def _db_file() -> Path:
    return ROOT / "data" / f"test-ingest-{uuid4().hex}.sqlite3"


@pytest.fixture(scope="module")
def seeded():
    """One seeded database for checks that save nothing."""
    path = _db_file()
    try:
        site_seed.build(path, AS_OF)
        yield path
    finally:
        path.unlink(missing_ok=True)


@pytest.fixture()
def db_path():
    path = _db_file()
    try:
        site_seed.build(path, AS_OF)
        yield path
    finally:
        path.unlink(missing_ok=True)


def xlsx(*rows, header=HEADER, above=(), sheet="readings", first_sheet=None) -> bytes:
    book = Workbook()
    if first_sheet:
        book.active.title = first_sheet
        target = book.create_sheet(sheet)
    else:
        target = book.active
        target.title = sheet
    for line in above:
        target.append(list(line))
    if header is not None:
        target.append(list(header))
    for row in rows:
        target.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def messages(content, db) -> list[str]:
    return [error["message"] for error in ingest_excel.preview(content, db, now=NOW)["errors"]]


def query(db, sql, *params) -> list[dict]:
    return site_db._all(sql, params, db)


def test_samples_are_marked_synthetic() -> None:
    for path in (SAMPLE, BAD_SAMPLE):
        book = load_workbook(path, read_only=True)
        try:
            assert "SYNTHETIC SAMPLE" in str(book["readings"]["A1"].value)
            assert "SYNTHETIC SAMPLE" in str(book["About"]["A1"].value)
        finally:
            book.close()


def test_sample_preview_commit_and_reimport(db_path) -> None:
    raw = SAMPLE.read_bytes()
    preview = ingest_excel.preview(SAMPLE, db_path, now=NOW)
    assert preview["errors"] == [] and preview["total_rows"] == len(preview["rows"]) == 162  # 3 meters x 9 times x 6 points
    assert {row["asset_id"] for row in preview["rows"]} == {"FDM-AHU-01", "FDM-PMP-01", "SGV-DB-01"}
    assert {row["status"] for row in preview["rows"]} == {"replaces_seed"}  # the seed already fills that window
    assert preview["rows"][0]["observed_at"] == "2026-10-09T00:00:00+00:00"  # 08:00 SGT stored as UTC
    assert not query(db_path, "SELECT * FROM ingest_batches WHERE source_type = 'excel'")  # a preview saves nothing

    result = ingest_excel.commit(SAMPLE, "electrical-sample.xlsx", db_path, now=NOW)
    assert result["batch_id"] == "excel-20261012-001"
    assert (result["saved"], result["skipped"], result["replaced_seed"], result["already_imported"]) == (162, 0, 162, False)
    batch = query(db_path, "SELECT * FROM ingest_batches WHERE batch_id = ?", result["batch_id"])[0]
    assert batch == {"batch_id": "excel-20261012-001", "source_type": "excel", "file_name": "electrical-sample.xlsx",
                     "sha256": hashlib.sha256(raw).hexdigest(), "received_at": "2026-10-12T06:07:00+00:00", "row_count": 162}
    assert query(db_path, "SELECT COUNT(*) AS n FROM readings WHERE ingest_batch_id = ?", batch["batch_id"])[0]["n"] == 162
    first = site_db.readings("FDM-AHU-01", "2026-10-09T00:00:00+00:00", "2026-10-09T00:00:00+00:00", db_path)
    sample_volts = next(row["value"] for row in preview["rows"]
                        if row["asset_id"] == "FDM-AHU-01" and row["point_id"] == "voltage_v")
    assert {(r["point_id"], r["ingest_batch_id"]) for r in first} >= {("voltage_v", "excel-20261012-001")}
    assert next(r["value"] for r in first if r["point_id"] == "voltage_v") == sample_volts

    again = ingest_excel.commit(SAMPLE, "electrical-sample.xlsx", db_path, now=NOW + timedelta(hours=1))
    assert (again["batch_id"], again["saved"], again["skipped"], again["already_imported"]) == ("excel-20261012-001", 0, 162, True)
    assert len(query(db_path, "SELECT * FROM ingest_batches WHERE source_type = 'excel'")) == 1

    # A different file whose rows are already present: duplicates skipped, recorded as the next batch of the day.
    overlap = xlsx(GOOD[:3] + (sample_volts, "V"))
    upload = io.BytesIO(overlap)
    upload.name = "re-export.xlsx"
    second = ingest_excel.commit(upload, None, db_path, now=NOW)
    assert (second["batch_id"], second["saved"], second["skipped"]) == ("excel-20261012-002", 0, 1)
    assert query(db_path, "SELECT file_name, row_count FROM ingest_batches WHERE batch_id = ?", second["batch_id"]) == [
        {"file_name": "re-export.xlsx", "row_count": 0}]


def test_imported_reading_survives_seed_rebuild(db_path) -> None:
    ingest_excel.commit(SAMPLE, "electrical-sample.xlsx", db_path, now=NOW)
    ingest_excel.commit(xlsx(("EMS-FDMAHU01", OLD, "current_a", 88.5, "A")), "old.xlsx", db_path, now=NOW)
    window = ("2026-10-09T00:00:00+00:00", "2026-10-09T02:00:00+00:00")
    before = site_db.readings("FDM-AHU-01", *window, db_path)
    site_seed.build(db_path, AS_OF + timedelta(days=1))
    after = site_db.readings("FDM-AHU-01", *window, db_path)
    assert after == before and {row["ingest_batch_id"] for row in after} == {"excel-20261012-001"}
    old = site_db.readings("FDM-AHU-01", "2026-09-01T00:00:00+00:00", "2026-09-01T00:00:00+00:00", db_path)
    assert [(row["value"], row["ingest_batch_id"]) for row in old] == [(88.5, "excel-20261012-002")]
    assert len(query(db_path, "SELECT * FROM ingest_batches WHERE source_type = 'seed'")) == 1
    assert ingest_excel.commit(SAMPLE, "electrical-sample.xlsx", db_path, now=NOW)["already_imported"]


def test_new_reading_and_sgt_batch_date(db_path) -> None:
    late = datetime(2026, 10, 12, 17, 0, tzinfo=timezone.utc)  # 01:00 SGT on 13 October
    content = xlsx(("EMS-FDMAHU01", OLD, "voltage_v", 231, "V"), ("ems-fdmahu01", "2026-09-01T00:15:00Z", "trip_state", 0, "state"))
    preview = ingest_excel.preview(content, db_path, now=late)
    assert [row["status"] for row in preview["rows"]] == ["new", "new"]
    result = ingest_excel.commit(content, "meters.xlsx", db_path, now=late)
    assert (result["batch_id"], result["saved"], result["replaced_seed"]) == ("excel-20261013-001", 2, 0)


def test_bad_sample_is_refused_whole(db_path) -> None:
    preview = ingest_excel.preview(BAD_SAMPLE, db_path, now=NOW)
    assert [error["row"] for error in preview["errors"]] == [22, 23, 24]
    assert "no UTC offset" in preview["errors"][0]["message"]
    assert "unknown meter_ref 'EMS-XYZ999'" in preview["errors"][1]["message"]
    assert "power_factor 1.4" in preview["errors"][2]["message"]
    assert len(preview["rows"]) == 18  # the valid rows still preview
    with pytest.raises(ingest_excel.ExcelImportError) as refused:
        ingest_excel.commit(BAD_SAMPLE, "electrical-bad-sample.xlsx", db_path, now=NOW)
    assert len(refused.value.preview["errors"]) == 3
    assert not query(db_path, "SELECT * FROM ingest_batches WHERE source_type = 'excel'")
    assert not query(db_path, "SELECT * FROM readings WHERE ingest_batch_id LIKE 'excel-%'")


@pytest.mark.parametrize("row, expected", [
    (("EMS-FDMAHU01", "2026-10-09T08:00:00", "voltage_v", 230, "V"), "has no UTC offset"),
    (("EMS-FDMAHU01", datetime(2026, 10, 9, 8, 0), "voltage_v", 230, "V"), "is an Excel date cell"),
    (("EMS-FDMAHU01", "09/10/2026 08:00", "voltage_v", 230, "V"), "is not an ISO 8601 time"),
    (("EMS-FDMAHU01", 46304.33, "voltage_v", 230, "V"), "timestamp is a number"),
    (("EMS-FDMAHU01", None, "voltage_v", 230, "V"), "timestamp is empty"),
    (("EMS-FDMAHU01", "2026-10-12T20:00:00+08:00", "voltage_v", 230, "V"), "is in the future"),
    (("EMS-FDMAHU01", "1999-12-31T08:00:00+08:00", "voltage_v", 230, "V"), "is before 2000"),
    (("EMS-NOPE01", "2026-10-09T08:00:00+08:00", "voltage_v", 230, "V"), "unknown meter_ref 'EMS-NOPE01'"),
    (("FDM-AHU-01", "2026-10-09T08:00:00+08:00", "voltage_v", 230, "V"), "unknown meter_ref"),  # asset IDs are not aliases
    ((None, "2026-10-09T08:00:00+08:00", "voltage_v", 230, "V"), "meter_ref is empty"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "frequency_hz", 50, "Hz"), "unknown point 'frequency_hz'"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", 0.23, "kV"), "unit for voltage_v must be 'V'"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "power_kw", 30, "KW"), "unit for power_kw must be 'kW'"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "power_factor", 0.9, "%"), "unit for power_factor must be empty"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "trip_state", 0, None), "unit for trip_state must be 'state'"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", "abc", "V"), "is not a plain number"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", "nan", "V"), "is not a plain number"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", None, "V"), "value is empty"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", True, "V"), "TRUE/FALSE"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "voltage_v", 400, "V"), "voltage_v 400 is outside the plausible range"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "current_a", -5, "A"), "current_a -5 is outside the plausible range"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "power_factor", 1.4, None), "power_factor 1.4 is outside"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "trip_state", 0.5, "state"), "trip_state must be 0 or 1"),
    (("EMS-FDMAHU01", "2026-10-09T08:00:00+08:00", "trip_state", 2, "state"), "trip_state must be 0 or 1"),
])
def test_each_row_rejection_reason(seeded, row, expected) -> None:
    preview = ingest_excel.preview(xlsx(GOOD, row), seeded, now=NOW)
    assert [error["row"] for error in preview["errors"]] == [3]  # Excel row numbers: header 1, good row 2, bad row 3
    assert expected in preview["errors"][0]["message"]
    assert len(preview["rows"]) == 1


def test_one_message_lists_every_problem_in_a_row(seeded) -> None:
    [message] = messages(xlsx(("EMS-NOPE01", "2026-10-09T08:00:00", "voltage_v", 999, "V")), seeded)
    assert "unknown meter_ref" in message and "no UTC offset" in message and "outside the plausible range" in message


@pytest.mark.parametrize("content, expected", [
    (b"", "The file is empty"),
    (b"meter_ref,timestamp,point,value,unit\n", "Not a readable .xlsx workbook"),  # a CSV renamed .xlsx
    (xlsx(GOOD, header=("meter_ref", "timestamp", "point", "value")), "No header row with meter_ref"),
    (xlsx(GOOD + ("x",), header=HEADER + ("unit",)), "Column(s) appear twice: unit"),
    (xlsx(GOOD + ("Bay South",), header=HEADER + ("site",)), "Unsupported column(s): site"),
    (xlsx(GOOD + (None, "stray")), "column 7 has a value but no heading"),
    (xlsx(), "No data rows under the header"),
    (xlsx(GOOD, GOOD[:3] + (231.0, "V")), "conflicts with row 2"),
])
def test_file_level_rejections(seeded, content, expected) -> None:
    assert any(expected in message for message in messages(content, seeded)), messages(content, seeded)


def test_layout_tolerance(seeded) -> None:
    """Title lines above the header, header case, a leading About sheet, Z offsets and blank rows are all fine."""
    content = xlsx(("EMS-FDMAHU01", "2026-10-09T00:00:00Z", "VOLTAGE_V", "230.4", "V"), (None,) * 5, GOOD,
                   header=("Meter_Ref", "Timestamp", "Point", "Value", "Unit"),
                   above=(("SYNTHETIC SAMPLE export",), ("Generated by a fictional EMS",)), first_sheet="About")
    preview = ingest_excel.preview(content, seeded, now=NOW)
    assert preview["errors"] == [] and preview["sheet"] == "readings" and preview["total_rows"] == 2
    assert [row["status"] for row in preview["rows"]] == ["replaces_seed", "duplicate_in_file"]  # same key, same value
    assert [row["row"] for row in preview["rows"]] == [4, 6]


def test_conflict_with_an_imported_reading_refuses_the_file(db_path) -> None:
    ingest_excel.commit(xlsx(("EMS-FDMAHU01", OLD, "voltage_v", 230, "V")), "first.xlsx", db_path, now=NOW)
    changed = xlsx(("EMS-FDMAHU01", OLD, "voltage_v", 232, "V"))
    assert messages(changed, db_path) == ["conflicts with a reading already saved (batch excel-20261012-001, value 230)"]
    with pytest.raises(ingest_excel.ExcelImportError):
        ingest_excel.commit(changed, "second.xlsx", db_path, now=NOW)
    assert len(query(db_path, "SELECT * FROM ingest_batches WHERE source_type = 'excel'")) == 1


def test_empty_database_says_build_it_first() -> None:
    path = _db_file()
    try:
        assert "no meter aliases" in messages(xlsx(GOOD), path)[0]
    finally:
        path.unlink(missing_ok=True)
