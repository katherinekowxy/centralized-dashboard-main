"""Electrical readings from an Excel (.xlsx) meter export into the site database (docs/site-data-sources.md).

The file
- One worksheet: the sheet named ``readings`` if there is one, otherwise the first sheet.
- A header row with exactly these columns, in any order: ``meter_ref``, ``timestamp``, ``point``, ``value``, ``unit``.
  The header may sit below a few title lines (the first 10 rows are searched); anything above it is ignored.
- One reading per row. Fully blank rows are skipped. Error row numbers are Excel's own row numbers.

The timestamp rule (one rule, no guessing)
- ``timestamp`` is ISO 8601 *text* with a UTC offset, for example ``2026-10-12T14:00:00+08:00`` (SGT) or
  ``2026-10-12T06:00:00Z``. It is stored as UTC.
- A naive timestamp (no offset) is refused, and so is an Excel date cell, because .xlsx date cells carry no time
  zone. A timestamp more than 15 minutes in the future, or before 2000, is refused: usually a wrong offset.

Each row
- ``meter_ref`` must be an alias in ``asset_aliases`` (for example ``EMS-FDMCHL01``); an unknown ref is refused.
- ``point`` and ``unit`` must be one of the pairs in ``POINTS``; ``power_factor`` has an empty unit.
- ``value`` must be a plain number inside the point's plausible range; ``trip_state`` is exactly 0 or 1.

Saving
- ``preview`` saves nothing. ``commit`` checks everything again inside one transaction and refuses the whole file if
  any row is bad (the same rule as ``demo/imports.py``).
- Duplicates are skipped: a reading already saved by an earlier file with the same value, or repeated in this file.
  The same file again (same SHA-256) returns its earlier batch and saves nothing new.
- A reading at the same asset, point and time as a *seed* reading replaces it, because the seed is a stand-in for
  real data (``site_seed.build`` likewise never overwrites an imported reading). A different value at the same key
  from another imported file is a conflict, and the file is refused.
- Every saved reading carries the batch ID ``excel-YYYYMMDD-NNN`` (SGT date of receipt, next free number), and the
  batch row records the file name, SHA-256, receipt time and the number of readings it wrote.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import sqlite3
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from demo import site_db

SGT = timezone(timedelta(hours=8), "SGT")
SOURCE_TYPE = "excel"
SHEET = "readings"
COLUMNS = ("meter_ref", "timestamp", "point", "value", "unit")
HEADER_SEARCH_ROWS = 10
MAX_FILE_BYTES = 20_000_000
MAX_ROWS = 100_000
EARLIEST = datetime(2000, 1, 1, tzinfo=timezone.utc)
FUTURE_TOLERANCE = timedelta(minutes=15)
_NUMBER = re.compile(r"[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?\Z")


@dataclass(frozen=True)
class Point:
    unit: str
    low: float
    high: float


POINTS = {  # plausible ranges for a 230 V (phase) site supply; outside them a reading is a data error, not a fault
    "voltage_v": Point("V", 0, 300),
    "current_a": Point("A", 0, 5_000),
    "power_kw": Point("kW", 0, 5_000),
    "energy_kwh": Point("kWh", 0, 1_000_000_000),
    "power_factor": Point("", 0, 1),
    "trip_state": Point("state", 0, 1),
}


class ExcelImportError(ValueError):
    """The file was refused; ``preview`` holds every row error. Nothing was saved."""

    def __init__(self, preview: dict[str, Any]):
        self.preview = preview
        super().__init__(f"Excel import refused with {len(preview['errors'])} error(s); nothing was saved")


# Helpers shared with demo/ingest_ptw_pdf.py, so both importers follow one time rule and one batch-numbering rule.

def read_source(source: Any) -> tuple[bytes, str | None]:
    """Bytes and file name (when known) from a path, raw bytes, or an uploaded file object (``getvalue``/``read``)."""
    if isinstance(source, (bytes, bytearray, memoryview)):
        return bytes(source), None
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_bytes(), path.name
    for method in ("getvalue", "read"):
        if hasattr(source, method):
            return bytes(getattr(source, method)()), getattr(source, "name", None)
    raise TypeError("source must be a file path, bytes, or a file object")


def aware_now(now: datetime | None = None) -> datetime:
    now = now or datetime.now(timezone.utc)
    return now if now.tzinfo else now.replace(tzinfo=timezone.utc)


def offset_time(value: Any, label: str = "timestamp") -> datetime:
    """ISO 8601 text with a UTC offset, as an aware datetime. Naive text and Excel date cells raise ValueError."""
    if isinstance(value, (datetime, date, time)):
        raise ValueError(f"{label} is an Excel date cell, which carries no time zone; enter it as text with an offset, "
                         "for example 2026-10-12T14:00:00+08:00")
    if isinstance(value, (int, float)):
        raise ValueError(f"{label} is a number; enter it as ISO text with an offset, for example 2026-10-12T14:00:00+08:00")
    text = "" if value is None else str(value).strip()
    if not text:
        raise ValueError(f"{label} is empty")
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        raise ValueError(f"{label} '{text}' is not an ISO 8601 time, for example 2026-10-12T14:00:00+08:00") from None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError(f"{label} '{text}' has no UTC offset; add one, for example +08:00 for SGT")
    return parsed


def next_batch_id(connection: sqlite3.Connection, source_type: str, received_at: datetime) -> str:
    """``<source>-YYYYMMDD-NNN``: the SGT date of receipt and the next free number that day."""
    prefix = f"{source_type}-{received_at.astimezone(SGT):%Y%m%d}-"
    used = [row[0][len(prefix):] for row in
            connection.execute("SELECT batch_id FROM ingest_batches WHERE batch_id LIKE ?", (prefix + "%",))]
    return f"{prefix}{max((int(n) for n in used if n.isdigit()), default=0) + 1:03d}"


def clean_file_name(name: Any, fallback: str) -> str:
    name = str(name or "").replace("\\", "/").rsplit("/", 1)[-1].strip() or fallback
    if len(name) > 255:
        raise ValueError("file name must be 255 characters or fewer")
    return name


# Row checks

def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float:
    if isinstance(value, bool):
        raise ValueError("value is TRUE/FALSE, not a number")
    if isinstance(value, (datetime, date, time)):
        raise ValueError("value is a date cell, not a number")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _text(value)
        if not text:
            raise ValueError("value is empty")
        if not _NUMBER.match(text):
            raise ValueError(f"value '{text}' is not a plain number")
        number = float(text)
    if not math.isfinite(number):
        raise ValueError("value is not a finite number")
    return number


def _check_value(point_id: str, point: Point, value: float) -> str | None:
    if point_id == "trip_state" and value not in (0.0, 1.0):
        return f"trip_state must be 0 or 1, not {value:g}"
    if not point.low <= value <= point.high:
        return f"{point_id} {value:g} is outside the plausible range {point.low:g} to {point.high:g}"
    return None


def _aliases(connection: sqlite3.Connection) -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    for alias, asset_id in connection.execute("SELECT alias, asset_id FROM asset_aliases"):
        out.setdefault(alias.strip().casefold(), set()).add(asset_id)
    return out


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """(Excel row number, column index by name) of the first row that names every column."""
    for number, row in enumerate(rows[:HEADER_SEARCH_ROWS], start=1):
        names = [_text(cell).lower() for cell in row]
        if set(COLUMNS) <= set(names):
            return number, {name: names.index(name) for name in COLUMNS}
    return None


def _evaluate(connection: sqlite3.Connection, raw: bytes, now: datetime) -> dict[str, Any]:
    result: dict[str, Any] = {"rows": [], "errors": [], "total_rows": 0, "sha256": hashlib.sha256(raw).hexdigest(),
                              "sheet": None, "new": 0, "duplicates": 0, "replaces_seed": 0}

    def refuse(row: int, message: str) -> dict[str, Any]:
        result["errors"].append({"row": row, "message": message})
        return result

    if not raw:
        return refuse(0, "The file is empty")
    if len(raw) > MAX_FILE_BYTES:
        return refuse(0, f"The file is larger than {MAX_FILE_BYTES // 1_000_000} MB")
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError, XML errors: all mean "not a usable .xlsx"
        return refuse(0, f"Not a readable .xlsx workbook ({type(exc).__name__}); save the export as Excel Workbook (.xlsx)")
    try:
        sheet = next((ws for ws in workbook.worksheets if ws.title.strip().lower() == SHEET), workbook.worksheets[0])
        result["sheet"] = sheet.title
        sheet.reset_dimensions()  # some exporters write a wrong size tag; read every cell actually present
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        return refuse(0, f"The worksheet could not be read ({type(exc).__name__})")
    finally:
        workbook.close()

    header = _find_header(rows)
    if header is None:
        return refuse(0, f"No header row with {', '.join(COLUMNS)} in the first {HEADER_SEARCH_ROWS} rows of sheet "
                         f"'{result['sheet']}'")
    header_row, index = header
    used = set(index.values())
    names = [_text(cell).lower() for cell in rows[header_row - 1]]
    repeated = sorted({name for name in COLUMNS if names.count(name) > 1})
    unexpected = sorted({name for name in names if name and name not in COLUMNS})
    if repeated:
        return refuse(header_row, f"Column(s) appear twice: {', '.join(repeated)}")
    if unexpected:
        return refuse(header_row, f"Unsupported column(s): {', '.join(unexpected)}; the export needs exactly "
                                  f"{', '.join(COLUMNS)}")
    aliases = _aliases(connection)
    if not aliases:
        return refuse(0, "The site database has no meter aliases yet; build it first (python -m demo.site_seed)")

    seen: dict[tuple[str, str, str], tuple[int, float]] = {}
    for number, row in enumerate(rows[header_row:], start=header_row + 1):
        if all(_text(cell) == "" for cell in row):
            continue
        result["total_rows"] += 1
        if result["total_rows"] > MAX_ROWS:
            refuse(number, f"More than {MAX_ROWS} data rows; split the export")
            break
        cell = {name: row[i] if i < len(row) else None for name, i in index.items()}
        problems = [f"column {i + 1} has a value but no heading" for i, value in enumerate(row) if i not in used and _text(value)]

        ref, asset_id = _text(cell["meter_ref"]), None
        owners = aliases.get(ref.casefold(), set())
        if not ref:
            problems.append("meter_ref is empty")
        elif not owners:
            problems.append(f"unknown meter_ref '{ref}' (no such alias in asset_aliases)")
        elif len(owners) > 1:
            problems.append(f"meter_ref '{ref}' matches more than one asset")
        else:
            asset_id = next(iter(owners))

        observed_at = None
        try:
            moment = offset_time(cell["timestamp"])
            if moment < EARLIEST:
                problems.append(f"timestamp {moment.isoformat()} is before 2000")
            elif moment > now + FUTURE_TOLERANCE:
                problems.append(f"timestamp {moment.isoformat()} is in the future; check the UTC offset")
            else:
                observed_at = site_db.utc(moment)
        except ValueError as exc:
            problems.append(str(exc))

        point_id, unit = _text(cell["point"]).lower(), _text(cell["unit"])
        point = POINTS.get(point_id)
        if point is None:
            problems.append(f"unknown point '{point_id}'; use one of {', '.join(POINTS)}")
        elif unit != point.unit:
            problems.append(f"unit for {point_id} must be " + (f"'{point.unit}'" if point.unit else "empty")
                            + (f", not '{unit}'" if unit else ""))

        value = None
        try:
            value = _number(cell["value"])
            if point is not None and (problem := _check_value(point_id, point, value)):
                problems.append(problem)
        except ValueError as exc:
            problems.append(str(exc))

        if not problems and asset_id and observed_at and point is not None and value is not None:
            key = (asset_id, point_id, observed_at)
            status = None
            if key in seen:
                first_row, first_value = seen[key]
                if math.isclose(first_value, value, rel_tol=1e-9, abs_tol=1e-9):
                    status = "duplicate_in_file"
                else:
                    problems.append(f"conflicts with row {first_row}: same meter, point and time, different value")
            else:
                existing = connection.execute(
                    """SELECT r.value, r.unit, r.ingest_batch_id, b.source_type FROM readings r
                       LEFT JOIN ingest_batches b ON b.batch_id = r.ingest_batch_id
                       WHERE r.asset_id = ? AND r.point_id = ? AND r.observed_at = ?""", key).fetchone()
                if existing is None:
                    status = "new"
                elif existing["source_type"] == "seed":
                    status = "replaces_seed"
                elif (existing["value"] is not None and existing["unit"] == point.unit
                      and math.isclose(existing["value"], value, rel_tol=1e-9, abs_tol=1e-9)):
                    status = "duplicate"
                else:
                    saved = "" if existing["value"] is None else f", value {existing['value']:g}"
                    problems.append(f"conflicts with a reading already saved (batch {existing['ingest_batch_id']}{saved})")
            if status:
                seen.setdefault(key, (number, value))
                result["rows"].append({"row": number, "meter_ref": ref, "asset_id": asset_id, "point_id": point_id,
                                       "observed_at": observed_at, "value": value, "unit": point.unit, "status": status})
                bucket = {"new": "new", "replaces_seed": "replaces_seed"}.get(status, "duplicates")
                result[bucket] += 1
        if problems:
            refuse(number, "; ".join(problems))

    if result["total_rows"] == 0 and not result["errors"]:
        refuse(header_row + 1, "No data rows under the header")
    return result


def preview(path_or_bytes: Any, db_path: Path | str | None = None, *, now: datetime | None = None) -> dict[str, Any]:
    """Check a file without saving anything.

    Returns ``rows`` (each valid row with its ``status``: new, replaces_seed, duplicate or duplicate_in_file),
    ``errors`` (``{"row": Excel row, "message": ...}``; row 0 means the whole file), ``total_rows``, ``sha256``,
    ``sheet`` and the counts ``new``, ``duplicates`` and ``replaces_seed``.
    """
    raw, _ = read_source(path_or_bytes)
    with closing(site_db.connect(db_path)) as connection:
        return _evaluate(connection, raw, aware_now(now))


def commit(path_or_bytes: Any, file_name: str | None = None, db_path: Path | str | None = None, *,
           now: datetime | None = None) -> dict[str, Any]:
    """Check again and save in one transaction, or raise ``ExcelImportError`` and save nothing.

    Returns ``batch_id``, ``saved`` (readings written, including seed readings replaced), ``skipped`` (duplicates),
    ``replaced_seed``, ``total_rows``, ``sha256`` and ``already_imported`` (the same file was saved before).
    """
    raw, path_name = read_source(path_or_bytes)
    name = clean_file_name(file_name or path_name, "upload.xlsx")
    now = aware_now(now)
    connection = site_db.connect(db_path)
    try:
        connection.execute("BEGIN IMMEDIATE")
        result = _evaluate(connection, raw, now)
        if result["errors"]:
            raise ExcelImportError(result)
        write = [row for row in result["rows"] if row["status"] in ("new", "replaces_seed")]
        earlier = connection.execute("SELECT batch_id FROM ingest_batches WHERE source_type = ? AND sha256 = ? "
                                     "ORDER BY batch_id LIMIT 1", (SOURCE_TYPE, result["sha256"])).fetchone()
        if earlier:
            batch_id = earlier["batch_id"]
            connection.execute("UPDATE ingest_batches SET row_count = row_count + ? WHERE batch_id = ?", (len(write), batch_id))
        else:
            batch_id = next_batch_id(connection, SOURCE_TYPE, now)
            site_db.insert(connection, "ingest_batches", [{"batch_id": batch_id, "source_type": SOURCE_TYPE, "file_name": name,
                                                           "sha256": result["sha256"], "received_at": site_db.utc(now),
                                                           "row_count": len(write)}])
        for status, mode in (("new", "insert"), ("replaces_seed", "replace")):
            site_db.insert(connection, "readings", [
                {"asset_id": row["asset_id"], "point_id": row["point_id"], "observed_at": row["observed_at"],
                 "value": row["value"], "unit": row["unit"], "ingest_batch_id": batch_id}
                for row in write if row["status"] == status], mode=mode)
        connection.commit()
    except BaseException:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {"batch_id": batch_id, "saved": len(write), "skipped": result["duplicates"],
            "replaced_seed": result["replaces_seed"], "total_rows": result["total_rows"], "sha256": result["sha256"],
            "already_imported": earlier is not None}


def main(argv: list[str] | None = None) -> int:
    """``python -m demo.ingest_excel FILE.xlsx [--db PATH] [--commit]``: preview, or save with ``--commit``."""
    import argparse

    parser = argparse.ArgumentParser(description="Check an electrical .xlsx export, and save it with --commit.")
    parser.add_argument("file", help="the .xlsx export")
    parser.add_argument("--db", help="site database (default data/site.sqlite3)")
    parser.add_argument("--commit", action="store_true", help="save the readings; without it nothing is saved")
    args = parser.parse_args(argv)
    if args.commit:
        try:
            saved = commit(args.file, None, args.db)
        except ExcelImportError as exc:
            result = exc.preview
        else:
            note = " (this file was imported before)" if saved["already_imported"] else ""
            print(f"Saved {saved['saved']} readings, skipped {saved['skipped']} duplicates, "
                  f"batch {saved['batch_id']}{note}")
            return 0
    else:
        result = preview(args.file, args.db)
    print(f"{result['total_rows']} data rows: {result['new']} new, {result['replaces_seed']} replace seed readings, "
          f"{result['duplicates']} duplicates, {len(result['errors'])} errors")
    for error in result["errors"]:
        print(f"  row {error['row']}: {error['message']}")
    if result["errors"]:
        print("Refused: fix every row above, then import the file again. Nothing was saved.")
    elif not args.commit:
        print("Preview only; add --commit to save.")
    return 1 if result["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
