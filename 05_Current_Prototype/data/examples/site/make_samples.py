"""Build the synthetic sample files for the two site importers. SYNTHETIC SAMPLE: every value here is invented.

    python data/examples/site/make_samples.py

Writes, next to this script:
- electrical-sample.xlsx       2 hours of 15-minute readings (09 Oct 2026, 08:00-10:00 SGT) for three seed assets,
                               by their energy-meter aliases EMS-FDMAHU01, EMS-FDMPMP01 and EMS-SGVDB01
- electrical-bad-sample.xlsx   the first quarter-hour of those readings plus three bad rows: a naive timestamp,
                               an unknown meter and a power factor of 1.4, so the importer refuses the file
- PTW-2026-0901.pdf            an electrical isolation of an asset plus a path closure (template SD-PTW-1)
- PTW-2026-0902.pdf            one physical restriction on Meadow Bridge (template SD-PTW-1)

Re-running gives byte-identical files, so a database that has imported them sees the same SHA-256 again.
Permit numbers are in the 09xx range so they never collide with the seed's 01xx permits.
"""

from __future__ import annotations

import io
import math
import random
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[2]  # 05_Current_Prototype
sys.path.insert(0, str(PROJECT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from demo.ingest_ptw_pdf import render_pdf  # noqa: E402

SGT = timezone(timedelta(hours=8))
FIXED = datetime(2026, 10, 12, 0, 0, 0)  # file metadata time, so the bytes never change
START = datetime(2026, 10, 9, 8, 0, tzinfo=SGT)
STEPS = 9  # 08:00 to 10:00 SGT inclusive, every 15 minutes
METERS = (("EMS-FDMAHU01", 45), ("EMS-FDMPMP01", 30), ("EMS-SGVDB01", 250))  # alias, rated current (A)
BANNER = "SYNTHETIC SAMPLE - fictional meter export for SignalDesk; not from any Gardens by the Bay system"
HEADER = ("meter_ref", "timestamp", "point", "value", "unit")
ABOUT = (
    "SYNTHETIC SAMPLE",
    "Fictional readings made by data/examples/site/make_samples.py for the SignalDesk demo.",
    "Not from any Gardens by the Bay system, meter or vendor.",
    "",
    "Layout read by demo/ingest_excel.py: sheet 'readings', header row meter_ref | timestamp | point | value | unit.",
    "timestamp is ISO text with a UTC offset, e.g. 2026-10-09T08:00:00+08:00. Date cells and naive times are refused.",
    "meter_ref is an alias in the asset_aliases table (EMS- plus the asset ID without hyphens).",
    "Points and units: voltage_v V, current_a A, power_kw kW, energy_kwh kWh, power_factor (no unit), trip_state state.",
)

PERMITS = (
    {"ptw_id": "PTW-2026-0901", "zone_code": "ACT", "contractor_id": "CTR-NWD", "title": "Active Garden outlet replacement",
     "work_type": "electrical", "start_at": "2026-10-14 09:00+08:00", "end_at": "2026-10-14 17:00+08:00",
     "impacts": [("electrical", "asset", "ACT-OUT-01", "isolated", "Outlets isolated", "-", "-"),
                 ("physical", "route", "Active Garden east path", "closed", "East path closed",
                  "2026-10-14 09:00+08:00", "2026-10-14 13:00+08:00")]},
    {"ptw_id": "PTW-2026-0902", "zone_code": "MDB", "contractor_id": "CTR-TMB", "title": "Meadow Bridge handrail repair",
     "work_type": "civil", "start_at": "2026-10-15 08:00+08:00", "end_at": "2026-10-16 18:00+08:00",
     "impacts": [("physical", "zone", "MDB", "restricted", "Single-file crossing only", "-", "-")]},
)


def reading_rows(seed: int = 901) -> list[list]:
    """Healthy, Operational-looking readings in export order (time, then meter)."""
    rng = random.Random(seed)
    energy = {ref: rng.uniform(20_000, 80_000) for ref, _ in METERS}
    rows = []
    for step in range(STEPS):
        at = (START + step * timedelta(minutes=15)).isoformat()
        for ref, rated in METERS:
            volts = round(230 + rng.gauss(0, 0.8), 1)
            amps = round(rated * (0.55 + rng.gauss(0, 0.02)), 1)
            pf = round(min(0.99, 0.91 + rng.gauss(0, 0.012)), 3)
            kw = round(math.sqrt(3) * volts * amps * pf / 1000, 2)  # the seed's formula, so charts stay consistent
            energy[ref] += kw * 0.25
            rows += [[ref, at, "voltage_v", volts, "V"], [ref, at, "current_a", amps, "A"], [ref, at, "power_kw", kw, "kW"],
                     [ref, at, "energy_kwh", round(energy[ref], 2), "kWh"], [ref, at, "power_factor", pf, None],
                     [ref, at, "trip_state", 0, "state"]]
    return rows


def bad_rows() -> list[list]:
    good = reading_rows()[: 6 * len(METERS)]  # the first quarter-hour, all valid
    return good + [
        ["EMS-FDMAHU01", "2026-10-09T08:15:00", "voltage_v", 230.4, "V"],                    # naive timestamp
        ["EMS-XYZ999", "2026-10-09T08:15:00+08:00", "current_a", 50.2, "A"],                 # unknown meter
        ["EMS-FDMPMP01", "2026-10-09T08:15:00+08:00", "power_factor", 1.4, None],           # power factor above 1
    ]


def workbook(rows: list[list], title: str) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = "readings"
    sheet["A1"] = BANNER
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="9E1C1C")
    sheet.append([])
    sheet.append(list(HEADER))
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column, width in zip("ABCDE", (16, 28, 14, 12, 8)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A4"
    about = book.create_sheet("About")
    for line in ABOUT:
        about.append([line])
    about["A1"].font = Font(bold=True, color="9E1C1C", size=14)
    about.column_dimensions["A"].width = 110
    book.properties.title = f"SYNTHETIC SAMPLE - {title}"
    book.properties.creator = "SignalDesk make_samples.py"
    book.properties.created = book.properties.modified = FIXED
    return book


def stable_bytes(book: Workbook) -> bytes:
    """The workbook as .xlsx bytes with fixed zip timestamps, so re-running gives the same SHA-256."""
    raw = io.BytesIO()
    book.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED.timetuple()[:6])
            entry.compress_type, entry.external_attr = zipfile.ZIP_DEFLATED, info.external_attr
            target.writestr(entry, source.read(info.filename))
    return out.getvalue()


def main() -> None:
    outputs = {
        "electrical-sample.xlsx": stable_bytes(workbook(reading_rows(), "electrical meter export")),
        "electrical-bad-sample.xlsx": stable_bytes(workbook(bad_rows(), "electrical meter export with three bad rows")),
        **{f"{permit['ptw_id']}.pdf": render_pdf(permit) for permit in PERMITS},
    }
    for name, data in outputs.items():
        (HERE / name).write_bytes(data)
        print(f"wrote {name} ({len(data):,} bytes)")


if __name__ == "__main__":
    main()
